import argparse
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook

from private_wire_workflow.companies_house import find_best_match, get_api_key


WORKBOOK_PATH = Path(
    "/Users/ssebl/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/"
    "5B98E93C-66D7-4736-914F-D6010E6BE16E/Target list origination (Alight).xlsx"
)
OUTPUT_PATH = Path("/Users/ssebl/Documents/New project/Target list origination (Alight) - CH enriched.xlsx")
SHEET_NAME = "Batch 2 account screening"


COLUMN_MAP = {
    "name": 1,
    "account_status": 3,
    "notes": 12,
    "rating_sources": 16,
    "revenue": 17,
    "bankability_confirmed": 18,
    "comment_bankability": 19,
}


def _bankability_comment(match) -> str:
    sic_text = ", ".join(match.sic_codes[:4]) if match.sic_codes else "n/a"
    return (
        f"Companies House match: {match.company_name} ({match.company_number}); "
        f"status={match.status or 'n/a'}; type={match.company_type or 'n/a'}; "
        f"incorporated={match.date_of_creation or 'n/a'}; "
        f"last accounts made up to={match.accounts_last_made_up_to or 'n/a'}; "
        f"next accounts due={match.accounts_next_due or 'n/a'}; "
        f"SIC={sic_text}; "
        "CH API does not provide a public credit rating or structured revenue figure, "
        "so rating/revenue still need external source or accounts-document review."
    )


def enrich_workbook(start_row: int = 2, max_rows: int = 100, save_every: int = 25) -> Path:
    api_key = get_api_key()
    if not OUTPUT_PATH.exists():
        copy2(WORKBOOK_PATH, OUTPUT_PATH)

    workbook = load_workbook(OUTPUT_PATH)
    sheet = workbook[SHEET_NAME]

    processed = 0
    for row in range(start_row, sheet.max_row + 1):
        if processed >= max_rows:
            break
        company_name = str(sheet.cell(row=row, column=COLUMN_MAP["name"]).value or "").strip()
        if not company_name:
            continue

        if str(sheet.cell(row=row, column=COLUMN_MAP["comment_bankability"]).value or "").strip():
            continue

        match = find_best_match(company_name, api_key)
        if match is None:
            sheet.cell(row=row, column=COLUMN_MAP["account_status"]).value = "Needs review"
            sheet.cell(row=row, column=COLUMN_MAP["notes"]).value = "No Companies House match found."
            sheet.cell(row=row, column=COLUMN_MAP["comment_bankability"]).value = (
                "No clear Companies House result. Manual legal-entity review needed."
            )
        else:
            sheet.cell(row=row, column=COLUMN_MAP["account_status"]).value = "Needs review"
            sheet.cell(row=row, column=COLUMN_MAP["notes"]).value = (
                f"Best CH match confidence={match.query_score:.2f}; "
                f"company number={match.company_number}; status={match.status}."
            )
            sheet.cell(row=row, column=COLUMN_MAP["rating_sources"]).value = match.source_url
            sheet.cell(row=row, column=COLUMN_MAP["bankability_confirmed"]).value = "Needs review"
            sheet.cell(row=row, column=COLUMN_MAP["comment_bankability"]).value = _bankability_comment(
                match
            )

        processed += 1
        if processed % save_every == 0:
            workbook.save(OUTPUT_PATH)
            print(f"Saved through row {row}", flush=True)

    workbook.save(OUTPUT_PATH)
    print(f"Saved final batch starting at row {start_row} with {processed} processed rows", flush=True)
    return OUTPUT_PATH


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100)
    parser.add_argument("--save-every", type=int, default=25)
    args = parser.parse_args()

    output = enrich_workbook(
        start_row=args.start_row,
        max_rows=args.max_rows,
        save_every=args.save_every,
    )
    print(output)
