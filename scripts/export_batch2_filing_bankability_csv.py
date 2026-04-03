import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from private_wire_workflow.companies_house import (
    download_document_content,
    find_best_match,
    get_api_key,
    list_accounts_filings,
    select_best_accounts_filing,
)
from private_wire_workflow.filing_bankability import assess_filing_bankability
from private_wire_workflow.filing_parser import (
    extract_financials_from_text,
    extract_text_from_pdf_bytes_with_metadata,
    html_to_text,
)


WORKBOOK_PATH = Path(
    "/Users/ssebl/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/"
    "5B98E93C-66D7-4736-914F-D6010E6BE16E/Target list origination (Alight).xlsx"
)
OUTPUT_PATH = Path("/Users/ssebl/Documents/New project/batch2_filing_bankability.csv")
SHEET_NAME = "Batch 2 account screening"


def _entity_scope(description: str) -> str:
    lowered = description.lower()
    if "group" in lowered:
        return "group"
    return "company"


def export_filing_bankability(start_row: int = 2, max_rows: int = 25) -> Path:
    api_key = get_api_key()
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]

    with OUTPUT_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "row_number",
                "input_name",
                "matched_company_name",
                "company_number",
                "filing_date",
                "filing_description",
                "entity_scope",
                "turnover",
                "ebitda",
                "interest_expense",
                "net_debt",
                "ebitda_margin",
                "interest_cover",
                "net_debt_to_ebitda",
                "verdict",
                "review_reason",
                "document_url",
                "source_format",
                "extraction_confidence",
                "extraction_engine",
            ]
        )

        processed = 0
        for row_number in range(start_row, sheet.max_row + 1):
            if processed >= max_rows:
                break

            company_name = str(sheet.cell(row=row_number, column=1).value or "").strip()
            if not company_name:
                continue

            match = find_best_match(company_name, api_key)
            if match is None:
                writer.writerow(
                    [
                        row_number,
                        company_name,
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "Needs review",
                        "No clear Companies House entity match.",
                        "",
                        "",
                        "",
                        "",
                    ]
                )
                processed += 1
                continue

            filings = list_accounts_filings(match.company_number, api_key)
            filing = select_best_accounts_filing(filings)
            if filing is None or not filing.document_metadata_url:
                writer.writerow(
                    [
                        row_number,
                        company_name,
                        match.company_name,
                        match.company_number,
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "Needs review",
                        "No usable accounts filing with document metadata found.",
                        "",
                        "",
                        "",
                        "",
                    ]
                )
                processed += 1
                continue

            document_bytes, source_format = download_document_content(
                filing.document_metadata_url,
                api_key,
            )
            if not document_bytes:
                writer.writerow(
                    [
                        row_number,
                        company_name,
                        match.company_name,
                        match.company_number,
                        filing.date,
                        filing.description,
                        _entity_scope(filing.description),
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "Needs review",
                        "Accounts document is unavailable for parsing.",
                        filing.document_metadata_url,
                        source_format,
                        "",
                        "",
                    ]
                )
                processed += 1
                continue

            if source_format == "pdf":
                text, engine, extract_notes = extract_text_from_pdf_bytes_with_metadata(
                    document_bytes
                )
                source_format = f"{source_format}:{engine}"
            else:
                text = html_to_text(document_bytes.decode("utf-8", errors="ignore"))
                extract_notes = []
            financials = extract_financials_from_text(text, source_format=source_format)
            financials.extraction_engine = source_format.split(":")[-1]
            financials.extraction_notes.extend(extract_notes)
            result = assess_filing_bankability(
                account_name=company_name,
                company_number=match.company_number,
                entity_name=match.company_name,
                filing=filing,
                financials=financials,
                document_url=filing.document_metadata_url,
                entity_scope=_entity_scope(filing.description),
            )

            writer.writerow(
                [
                    row_number,
                    company_name,
                    match.company_name,
                    match.company_number,
                    result.filing_date,
                    result.filing_description,
                    result.entity_scope,
                    result.turnover,
                    result.ebitda,
                    result.interest_expense,
                    result.net_debt,
                    result.ebitda_margin,
                    result.interest_cover,
                    result.net_debt_to_ebitda,
                    result.verdict,
                    result.review_reason,
                    result.document_url,
                    result.source_format,
                    financials.extraction_confidence,
                    financials.extraction_engine,
                ]
            )
            processed += 1
            if processed % 10 == 0:
                print(f"Exported {processed} filing rows through worksheet row {row_number}", flush=True)
    return OUTPUT_PATH


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=25)
    args = parser.parse_args()
    print(export_filing_bankability(start_row=args.start_row, max_rows=args.max_rows))
