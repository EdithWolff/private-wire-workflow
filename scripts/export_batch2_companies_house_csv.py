import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from private_wire_workflow.companies_house import find_best_match, get_api_key


WORKBOOK_PATH = Path(
    "/Users/ssebl/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/"
    "5B98E93C-66D7-4736-914F-D6010E6BE16E/Target list origination (Alight).xlsx"
)
OUTPUT_PATH = Path("/Users/ssebl/Documents/New project/batch2_companies_house_enrichment.csv")
SHEET_NAME = "Batch 2 account screening"


def export_matches(start_row: int = 2, max_rows: int = 100) -> Path:
    api_key = get_api_key()
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]

    with OUTPUT_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "row_number",
                "input_name",
                "matched_company_name",
                "company_number",
                "company_status",
                "company_type",
                "date_of_creation",
                "accounts_last_made_up_to",
                "accounts_next_due",
                "sic_codes",
                "source_url",
                "query_score",
                "match_quality",
                "bankability_comment",
            ]
        )

        processed = 0
        for row_number in range(start_row, sheet.max_row + 1):
            if processed >= max_rows:
                break

            company_name = str(sheet.cell(row=row_number, column=1).value or "").strip()
            if not company_name:
                continue

            match = find_best_match(company_name, api_key)
            match_quality = "Low"
            if match and match.query_score >= 0.95:
                match_quality = "High"
            elif match and match.query_score >= 0.80:
                match_quality = "Medium"

            if match is None:
                writer.writerow(
                    [
                        row_number,
                        company_name,
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                        "Low",
                        "No clear Companies House result. Manual legal-entity review needed.",
                    ]
                )
            else:
                writer.writerow(
                    [
                        row_number,
                        company_name,
                        match.company_name,
                        match.company_number,
                        match.status,
                        match.company_type,
                        match.date_of_creation,
                        match.accounts_last_made_up_to,
                        match.accounts_next_due,
                        ", ".join(match.sic_codes),
                        match.source_url,
                        f"{match.query_score:.2f}",
                        match_quality,
                        (
                            f"Companies House match: {match.company_name} ({match.company_number}); "
                            f"status={match.status or 'n/a'}; type={match.company_type or 'n/a'}; "
                            f"incorporated={match.date_of_creation or 'n/a'}; "
                            f"last accounts made up to={match.accounts_last_made_up_to or 'n/a'}; "
                            f"next accounts due={match.accounts_next_due or 'n/a'}."
                        ),
                    ]
                )

            processed += 1
            if processed % 25 == 0:
                print(f"Exported {processed} rows through worksheet row {row_number}", flush=True)

    return OUTPUT_PATH


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--max-rows", type=int, default=100)
    args = parser.parse_args()
    print(export_matches(start_row=args.start_row, max_rows=args.max_rows))
