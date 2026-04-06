import argparse
import csv
import multiprocessing as mp
import re
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

from private_wire_workflow.bankability_master import rating_at_least_bbb_minus
from private_wire_workflow.companies_house import (
    download_document_pdf_content,
    find_best_match,
    get_api_key,
    list_accounts_filings,
)
from private_wire_workflow.filing_text_assessment import (
    assess_previous_ppa,
    assess_sustainability_targets,
    compute_ratio_assessment_from_text,
    normalize_company_name,
    ocr_all_pdf_pages_with_tesseract,
    select_latest_non_dormant_full_accounts,
)
from private_wire_workflow.ratings import MOODYS_TO_SP, RATING_ORDER
from private_wire_workflow.report_utils import NO_INFO, fmt_number as _fmt_number, fmt_pct as _fmt_pct
DEFAULT_COMPANY_TIMEOUT_SEC = 900

BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_RATINGS_CSV = BASE_DIR / "data" / "inputs" / "company_rating_screening_findings.csv"
TXT_DIR = BASE_DIR / "data" / "filing_texts" / "pharma"
OUTPUT_XLSX = BASE_DIR / "data" / "pharma_60_case_by_case.xlsx"
RUN_SUMMARY_TXT = BASE_DIR / "data" / "pharma_60_run_summary.txt"

COMPANIES_60 = [
    "AbbVie Inc.",
    "Aldi UK",
    "Almac Group",
    "Almirall S.A.",
    "AMRI",
    "Asda Supermarket",
    "AstraZeneca",
    "Boehringer Ingelheim",
    "Boots UK",
    "Bristol-Myers Squibb Co",
    "Cambrex",
    "Elanco Animal Health Inc.",
    "Eli Lilly And Company",
    "Eurospar",
    "Galderma SA",
    "Giant Tiger",
    "Grifols S.A.",
    "GSK",
    "H-E-B",
    "Hikma Pharmaceuticals PLC",
    "Indivior plc",
    "Ipsen Pharma",
    "Jazz Pharmaceuticals PLC",
    "Lidl GB",
    "Lidl In Germany",
    "Lloyd's Pharmacy",
    "Lonza Group Ltd",
    "Mallinckrodt Pharmaceuticals",
    "McColl S Retail Group",
    "Melaleuca The Wellness Company",
    "Merck & Co. Inc.",
    "Merck KGaA",
    "Merz Aesthetics EMEA",
    "MSD",
    "Mundipharma International Limited",
    "Novartis AG",
    "One Stop Stores Ltd",
    "Organon & Co.",
    "Patheon",
    "Perrigo Company plc",
    "Pfizer Inc.",
    "Pick N Pay",
    "Pierre Fabre Group",
    "Recordati Industria Chimica E Farmaceutica S.p.A",
    "Rowlands Pharmacy",
    "Sandoz AG",
    "Sanofi",
    "Savers Health Home & Beauty",
    "Servier",
    "Shoppers Drug Mart",
    "STERIGENE",
    "Superdrug",
    "Takeda Pharmaceutical",
    "Tesco",
    "Teva Pharmaceutical Industries Ltd.",
    "The Janssen Pharmaceutical Companies Of Johnson & Johnson",
    "Viatris Inc.",
    "Virbac SA",
    "Walgreens Boots Alliance",
    "West Pharmaceutical Services",
]

ALIASES = {
    "teva pharmaceutical industries ltd": "teva pharmaceutical industries",
    "recordati industria chimica e farmaceutica spa": "recordati industria chimica",
    "the janssen pharmaceutical companies of johnson johnson": "the janssen pharmaceutical",
}


def _norm(value: str) -> str:
    text = "".join(ch.lower() for ch in value if ch.isalnum() or ch.isspace()).strip()
    return ALIASES.get(text, text)


def _extract_sp_like(text: str) -> str:
    if not text:
        return ""
    upper = text.upper()
    ordered = sorted(RATING_ORDER, key=len, reverse=True)
    for rating in ordered:
        pattern = rf"(?<![A-Z0-9]){re.escape(rating)}(?![A-Z0-9])"
        if re.search(pattern, upper):
            return rating
    return ""


def _extract_moodys(text: str) -> str:
    if not text:
        return ""
    m = re.search(
        r"\b(?:AAA|AA1|AA2|AA3|A1|A2|A3|BAA1|BAA2|BAA3|BA1|BA2|BA3|B1|B2|B3|CAA1|CAA2|CAA3|CAA|CA|C)\b",
        text.upper(),
    )
    return m.group(0) if m else ""


def _parse_ratings(combined: str) -> Tuple[str, str, str]:
    moodys = ""
    sp = ""
    fitch = ""
    chunks = [chunk.strip() for chunk in (combined or "").split("/") if chunk.strip()]
    for chunk in chunks:
        low = chunk.lower()
        if "mood" in low and not moodys:
            moodys = _extract_moodys(chunk) or moodys
            continue
        if "fitch" in low and not fitch:
            fitch = _extract_sp_like(chunk) or fitch
            continue
        parsed_sp = _extract_sp_like(chunk)
        if parsed_sp and not sp:
            sp = parsed_sp
            continue
        if not moodys:
            moodys = _extract_moodys(chunk) or moodys
    return moodys, sp, fitch


def _sp_equivalent(moodys: str, sp: str, fitch: str) -> str:
    if sp:
        return _extract_sp_like(sp) or sp
    if fitch:
        return _extract_sp_like(fitch) or fitch
    if moodys:
        return MOODYS_TO_SP.get(_extract_moodys(moodys), "")
    return ""


def _flags(metrics: Dict[str, Optional[float]]) -> Tuple[str, str, str]:
    margin = metrics.get("ebitda_margin")
    coverage = metrics.get("interest_coverage")
    leverage = metrics.get("net_debt_to_ebitda")
    margin_flag = "Yes" if margin is not None and margin >= 0.10 else ("No" if margin is not None else NO_INFO)
    coverage_flag = "Yes" if coverage is not None and coverage >= 3 else ("No" if coverage is not None else NO_INFO)
    leverage_flag = "Yes" if leverage is not None and leverage <= 2 else ("No" if leverage is not None else NO_INFO)
    return margin_flag, coverage_flag, leverage_flag


def _ratio_decision(metrics: Dict[str, Optional[float]]) -> Tuple[str, str]:
    margin = metrics.get("ebitda_margin")
    coverage = metrics.get("interest_coverage")
    leverage = metrics.get("net_debt_to_ebitda")
    if margin is None or coverage is None or leverage is None:
        return "Needs review", "Ratio-based decision: insufficient financial fields from filing OCR."
    if margin >= 0.10 and coverage >= 3 and leverage <= 2:
        return "Qualified", "Ratio-based decision: all three thresholds passed."
    return "Disqualified", "Ratio-based decision: one or more thresholds failed."


def _ppa_type_year(text: str) -> Tuple[str, str]:
    low = text.lower()
    has_ppa = assess_previous_ppa(text) == "Yes"
    if not has_ppa:
        return NO_INFO, NO_INFO
    wind = "wind" in low and "ppa" in low
    solar = "solar" in low and "ppa" in low
    if wind and solar:
        ptype = "Wind/Solar"
    elif wind:
        ptype = "Wind"
    elif solar:
        ptype = "Solar"
    else:
        ptype = NO_INFO
    year_m = re.search(r"\b(20\d{2})\b", text)
    return ptype, (year_m.group(1) if year_m else NO_INFO)


def _load_rating_map() -> Dict[str, Dict[str, str]]:
    rating_map: Dict[str, Dict[str, str]] = {}
    with INPUT_RATINGS_CSV.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            name = (row.get("Company Name") or "").strip()
            if not name:
                continue
            combined = (row.get("Credit Rating (S&P / Moody's / Fitch)") or "").strip()
            moodys, sp, fitch = _parse_ratings(combined)
            rating_map[_norm(name)] = {
                "moodys": moodys or NO_INFO,
                "sp": sp or NO_INFO,
                "fitch": fitch or NO_INFO,
            }
    return rating_map


def _headers() -> List[str]:
    return [
        "Company Name",
        "Run Status",
        "Comment",
        "Moody's Rating",
        "S&P Rating",
        "Turnover",
        "Fitch Rating",
        "EBITDA",
        "Interest Expense",
        "Debt",
        "Cash",
        "Net Debt",
        "EBITDA Margin",
        "Interest Coverage",
        "Net Debt/EBITDA",
        "Margin>=10%",
        "Coverage>=3x",
        "NetDebt/EBITDA<=2x",
        "Bankability Decision",
        "Disqualification Reason",
        "Renewable/Sustainability Signal",
        "Solar/PPA Signal",
        "PPA Type",
        "PPA Year",
    ]


def _blank_output(company_name: str, moodys: str, sp: str, fitch: str) -> Dict[str, str]:
    return {
        "Company Name": company_name,
        "Run Status": "failed_after_retries",
        "Comment": "Failed after retries.",
        "Moody's Rating": moodys,
        "S&P Rating": sp,
        "Turnover": NO_INFO,
        "Fitch Rating": fitch,
        "EBITDA": NO_INFO,
        "Interest Expense": NO_INFO,
        "Debt": NO_INFO,
        "Cash": NO_INFO,
        "Net Debt": NO_INFO,
        "EBITDA Margin": NO_INFO,
        "Interest Coverage": NO_INFO,
        "Net Debt/EBITDA": NO_INFO,
        "Margin>=10%": NO_INFO,
        "Coverage>=3x": NO_INFO,
        "NetDebt/EBITDA<=2x": NO_INFO,
        "Bankability Decision": "Needs review",
        "Disqualification Reason": NO_INFO,
        "Renewable/Sustainability Signal": NO_INFO,
        "Solar/PPA Signal": NO_INFO,
        "PPA Type": NO_INFO,
        "PPA Year": NO_INFO,
    }


def _load_existing_outputs(headers: List[str]) -> Dict[str, Dict[str, str]]:
    if not OUTPUT_XLSX.exists():
        return {}
    wb = load_workbook(OUTPUT_XLSX, data_only=True)
    if "Case_by_Case" not in wb.sheetnames:
        return {}
    ws = wb["Case_by_Case"]
    file_headers = [cell.value for cell in ws[1]]
    index = {h: i for i, h in enumerate(file_headers)}
    if "Company Name" not in index:
        return {}
    existing: Dict[str, Dict[str, str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        company_name = row[index["Company Name"]]
        if not company_name:
            continue
        data: Dict[str, str] = {}
        for h in headers:
            value = row[index[h]] if h in index and index[h] < len(row) else NO_INFO
            data[h] = NO_INFO if value in (None, "") else str(value)
        existing[str(company_name)] = data
    return existing


def _write_output_workbook(headers: List[str], rows: List[Dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Case_by_Case"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, NO_INFO) for h in headers])
    wb.save(OUTPUT_XLSX)


def _write_run_summary(headers: List[str], rows: List[Dict[str, str]]) -> None:
    status_counts: Dict[str, int] = {}
    for row in rows:
        status = row.get("Run Status", NO_INFO)
        status_counts[status] = status_counts.get(status, 0) + 1
    missing_txt: List[str] = []
    for row in rows:
        if row.get("Run Status") != "completed":
            continue
        company_name = row.get("Company Name", "")
        norm_name = normalize_company_name(company_name)
        found = any(path.name.startswith(f"{norm_name}__") for path in TXT_DIR.glob("*.txt"))
        if not found:
            missing_txt.append(company_name)
    failure_counts: Dict[str, int] = {}
    for row in rows:
        if row.get("Run Status") != "failed_after_retries":
            continue
        comment = row.get("Comment", NO_INFO)
        failure_counts[comment] = failure_counts.get(comment, 0) + 1
    top_failures = sorted(failure_counts.items(), key=lambda x: x[1], reverse=True)[:5]
    lines = [
        "pharma_60 run summary",
        f"rows_total={len(rows)}",
        f"txt_files_total={len(list(TXT_DIR.glob('*.txt')))}",
        "status_counts:",
    ]
    for status in sorted(status_counts):
        lines.append(f"- {status}: {status_counts[status]}")
    lines.append("completed_without_txt:")
    if missing_txt:
        lines.extend([f"- {name}" for name in missing_txt])
    else:
        lines.append("- none")
    lines.append("top_failure_reasons:")
    if top_failures:
        for reason, count in top_failures:
            lines.append(f"- {count}x {reason}")
    else:
        lines.append("- none")
    RUN_SUMMARY_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _ocr_worker(pdf_bytes: bytes, page_timeout_sec: int, queue: mp.Queue) -> None:
    try:
        ocr = ocr_all_pdf_pages_with_tesseract(pdf_bytes, page_timeout_sec=page_timeout_sec)
        queue.put(("ok", ocr.text))
    except Exception as exc:  # pragma: no cover - defensive boundary
        queue.put(("err", str(exc)))


def _ocr_text_with_timeout(pdf_bytes: bytes, page_timeout_sec: int, company_timeout_sec: int) -> str:
    queue: mp.Queue = mp.Queue()
    proc = mp.Process(target=_ocr_worker, args=(pdf_bytes, page_timeout_sec, queue), daemon=True)
    proc.start()
    proc.join(company_timeout_sec)
    if proc.is_alive():
        proc.terminate()
        proc.join(2)
        raise TimeoutError(f"OCR timed out after {company_timeout_sec}s")
    if queue.empty():
        raise RuntimeError("OCR worker exited without result")
    status, payload = queue.get()
    if status == "ok":
        return payload
    raise RuntimeError(payload)


def _no_filing_reason(filings: List) -> str:
    if not filings:
        return "No accounts-category filings returned for matched entity."
    with_metadata = [f for f in filings if f.document_metadata_url]
    if not with_metadata:
        return "Accounts filings exist but none expose document metadata."
    descriptions = [(f.description or "").lower() for f in with_metadata]
    if descriptions and all("dormant" in d for d in descriptions):
        return "Only dormant accounts filings were available."
    if descriptions and all("micro" in d for d in descriptions):
        return "Only micro-entity accounts filings were available."
    if descriptions and all(("full" not in d and "group" not in d) for d in descriptions):
        return "No full/group-tagged non-dormant accounts filing was available; fallback selection also failed."
    return "Matched entity found, but no qualifying non-dormant accounts filing could be selected."


def run(
    retries: int = 3,
    retry_backoff_sec: float = 2.0,
    ocr_page_timeout_sec: int = 30,
    company_timeout_sec: int = DEFAULT_COMPANY_TIMEOUT_SEC,
) -> None:
    api_key = get_api_key()
    TXT_DIR.mkdir(parents=True, exist_ok=True)
    rating_map = _load_rating_map()
    headers = _headers()
    existing = _load_existing_outputs(headers)
    reusable_terminal_statuses = {"completed", "no_match", "no_filing"}
    outputs: List[Dict[str, str]] = []

    for idx, company_name in enumerate(COMPANIES_60, start=1):
        ratings = rating_map.get(_norm(company_name), {"moodys": NO_INFO, "sp": NO_INFO, "fitch": NO_INFO})
        moodys = ratings["moodys"]
        sp = ratings["sp"]
        fitch = ratings["fitch"]
        existing_row = existing.get(company_name)
        if existing_row and existing_row.get("Run Status") in reusable_terminal_statuses:
            outputs.append(existing_row)
            print(f"[{idx}/60] {company_name}: reused_{existing_row.get('Run Status')}", flush=True)
            continue
        sp_eq = _sp_equivalent(
            moodys if moodys != NO_INFO else "",
            sp if sp != NO_INFO else "",
            fitch if fitch != NO_INFO else "",
        )
        output = _blank_output(company_name, moodys, sp, fitch)

        for attempt in range(1, retries + 1):
            try:
                match = find_best_match(company_name, api_key)
                if not match:
                    output["Run Status"] = "no_match"
                    output["Comment"] = (
                        "No confident UK legal entity match found after exact+normalized alias search."
                    )
                    if sp_eq:
                        decision = "Qualified" if rating_at_least_bbb_minus(sp_eq) else "Disqualified"
                        output["Bankability Decision"] = decision
                        output["Disqualification Reason"] = (
                            NO_INFO if decision == "Qualified" else "Credit rating below BBB- threshold."
                        )
                    break

                filings = list_accounts_filings(match.company_number, api_key)
                filing = select_latest_non_dormant_full_accounts(filings)
                if not filing:
                    output["Run Status"] = "no_filing"
                    output["Comment"] = _no_filing_reason(filings)
                    if sp_eq:
                        decision = "Qualified" if rating_at_least_bbb_minus(sp_eq) else "Disqualified"
                        output["Bankability Decision"] = decision
                        output["Disqualification Reason"] = (
                            NO_INFO if decision == "Qualified" else "Credit rating below BBB- threshold."
                        )
                    break

                pdf_bytes = download_document_pdf_content(filing.document_metadata_url, api_key)
                if not pdf_bytes:
                    output["Run Status"] = "no_filing"
                    output["Comment"] = (
                        "Matched entity and filing found, but filing PDF content was unavailable from document API."
                    )
                    if sp_eq:
                        decision = "Qualified" if rating_at_least_bbb_minus(sp_eq) else "Disqualified"
                        output["Bankability Decision"] = decision
                        output["Disqualification Reason"] = (
                            NO_INFO if decision == "Qualified" else "Credit rating below BBB- threshold."
                        )
                    break

                filing_year = (filing.date or "unknown")[:4]
                txt_name = f"{normalize_company_name(company_name)}__{filing_year}__{match.company_number}.txt"
                txt_path = TXT_DIR / txt_name
                if txt_path.exists() and txt_path.stat().st_size > 0:
                    ocr_text = txt_path.read_text(encoding="utf-8", errors="ignore")
                else:
                    ocr_text = _ocr_text_with_timeout(
                        pdf_bytes, page_timeout_sec=ocr_page_timeout_sec, company_timeout_sec=company_timeout_sec
                    )
                    txt_path.write_text(ocr_text, encoding="utf-8")

                metrics = compute_ratio_assessment_from_text(ocr_text)
                margin_flag, coverage_flag, leverage_flag = _flags(metrics)
                ratio_decision, ratio_comment = _ratio_decision(metrics)

                ppa_signal = assess_previous_ppa(ocr_text)
                ppa_type, ppa_year = _ppa_type_year(ocr_text)

                if sp_eq:
                    decision = "Qualified" if rating_at_least_bbb_minus(sp_eq) else "Disqualified"
                    comment = "Decision based on official rating (rating-first policy)."
                    disq = NO_INFO if decision == "Qualified" else "Credit rating below BBB- threshold."
                else:
                    decision = ratio_decision
                    comment = ratio_comment
                    disq = NO_INFO if decision != "Disqualified" else "Ratio threshold not met."

                output.update(
                    {
                        "Run Status": "completed",
                        "Comment": comment,
                        "Turnover": _fmt_number(metrics.get("turnover")),
                        "EBITDA": _fmt_number(metrics.get("ebitda")),
                        "Interest Expense": _fmt_number(metrics.get("interest_expense")),
                        "Debt": _fmt_number(metrics.get("debt")),
                        "Cash": _fmt_number(metrics.get("cash")),
                        "Net Debt": _fmt_number(metrics.get("net_debt")),
                        "EBITDA Margin": _fmt_pct(metrics.get("ebitda_margin")),
                        "Interest Coverage": _fmt_number(metrics.get("interest_coverage")),
                        "Net Debt/EBITDA": _fmt_number(metrics.get("net_debt_to_ebitda")),
                        "Margin>=10%": margin_flag,
                        "Coverage>=3x": coverage_flag,
                        "NetDebt/EBITDA<=2x": leverage_flag,
                        "Bankability Decision": decision,
                        "Disqualification Reason": disq,
                        "Renewable/Sustainability Signal": assess_sustainability_targets(ocr_text),
                        "Solar/PPA Signal": ppa_signal,
                        "PPA Type": ppa_type,
                        "PPA Year": ppa_year,
                    }
                )
                break
            except Exception as exc:
                output["Run Status"] = "failed_after_retries"
                output["Comment"] = f"Technical/API/OCR failure: {exc}"
                output["Bankability Decision"] = "Needs review"
                output["Disqualification Reason"] = NO_INFO
                if attempt >= retries:
                    break
                time.sleep(retry_backoff_sec * attempt)

        outputs.append(output)
        _write_output_workbook(headers, outputs)
        print(f"[{idx}/60] {company_name}: {output['Run Status']}", flush=True)
    _write_output_workbook(headers, outputs)
    _write_run_summary(headers, outputs)
    print(f"Saved: {OUTPUT_XLSX}", flush=True)
    print(f"TXT folder: {TXT_DIR}", flush=True)
    print(f"Run summary: {RUN_SUMMARY_TXT}", flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--retry-backoff-sec", type=float, default=2.0)
    parser.add_argument("--ocr-page-timeout-sec", type=int, default=30)
    parser.add_argument("--company-timeout-sec", type=int, default=DEFAULT_COMPANY_TIMEOUT_SEC)
    args = parser.parse_args()
    run(
        retries=args.retries,
        retry_backoff_sec=args.retry_backoff_sec,
        ocr_page_timeout_sec=args.ocr_page_timeout_sec,
        company_timeout_sec=args.company_timeout_sec,
    )
